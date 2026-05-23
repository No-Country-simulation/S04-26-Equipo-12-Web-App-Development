import io
from datetime import datetime

from django.db.models import Avg, Count, ExpressionWrapper, fields, Value
from django.db.models.functions import Coalesce, Extract, TruncDate
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.incidents.models import Incident, IncidentAssignment


def _parse_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.strptime(value, '%Y-%m-%d')
    except (ValueError, TypeError):
        try:
            return datetime.fromisoformat(value)
        except (ValueError, TypeError):
            return None


def get_avg_response_time(start_date, end_date):
    filters = {'is_current': True}
    sd = _parse_date(start_date)
    ed = _parse_date(end_date)
    if sd:
        filters['incident__created_at__gte'] = sd
    if ed:
        filters['incident__created_at__lte'] = ed

    assignments = (
        IncidentAssignment.objects
        .filter(**filters)
        .select_related('incident', 'incident__area')
        .annotate(
            response_minutes=ExpressionWrapper(
                (Extract('assigned_at', 'epoch') - Extract('incident__created_at', 'epoch')) / 60.0,
                output_field=fields.FloatField(),
            )
        )
        .values('incident__area__name')
        .annotate(
            avg_minutes=Coalesce(Avg('response_minutes'), Value(0.0)),
        )
        .order_by('incident__area__name')
    )
    return [
        {'area_name': a['incident__area__name'], 'avg_minutes': round(a['avg_minutes'], 2)}
        for a in assignments
    ]


def get_avg_resolution_time(start_date, end_date):
    filters = {'status': Incident.Status.CLOSED}
    sd = _parse_date(start_date)
    ed = _parse_date(end_date)
    if sd:
        filters['resolved_at__gte'] = sd
    if ed:
        filters['resolved_at__lte'] = ed

    incidents = (
        Incident.objects
        .filter(**filters)
        .values('area__name')
        .annotate(
            avg_minutes=Coalesce(
                Avg(
                    ExpressionWrapper(
                        (Extract('resolved_at', 'epoch') - Extract('created_at', 'epoch')) / 60.0,
                        output_field=fields.FloatField(),
                    )
                ),
                Value(0.0),
            ),
        )
        .order_by('area__name')
    )
    return [
        {'area_name': i['area__name'], 'avg_minutes': round(i['avg_minutes'], 2)}
        for i in incidents
    ]


def get_resolution_rate(start_date, end_date):
    filters = {}
    sd = _parse_date(start_date)
    ed = _parse_date(end_date)
    if sd:
        filters['created_at__gte'] = sd
    if ed:
        filters['created_at__lte'] = ed

    total = Incident.objects.filter(**filters).count()
    closed = Incident.objects.filter(**filters, status=Incident.Status.CLOSED).count()
    rate = round((closed / total * 100), 2) if total > 0 else 0.0
    return {'total': total, 'closed': closed, 'rate': rate}


def get_incidents_by_area(start_date, end_date):
    filters = {}
    sd = _parse_date(start_date)
    ed = _parse_date(end_date)
    if sd:
        filters['created_at__gte'] = sd
    if ed:
        filters['created_at__lte'] = ed

    qs = (
        Incident.objects
        .filter(**filters)
        .values('area__name')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    return [{'area_name': i['area__name'], 'count': i['count']} for i in qs]


def get_root_cause_frequency(start_date, end_date):
    filters = {'status': Incident.Status.CLOSED}
    sd = _parse_date(start_date)
    ed = _parse_date(end_date)
    if sd:
        filters['resolved_at__gte'] = sd
    if ed:
        filters['resolved_at__lte'] = ed

    qs = (
        Incident.objects
        .filter(**filters)
        .exclude(root_cause='')
        .values('type__name', 'root_cause')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    return [
        {'type_name': i['type__name'], 'root_cause': i['root_cause'], 'count': i['count']}
        for i in qs
    ]


def get_critical_incidents_over_time(start_date, end_date):
    filters = {'priority__in': [Incident.Priority.CRITICAL, Incident.Priority.HIGH]}
    sd = _parse_date(start_date)
    ed = _parse_date(end_date)
    if sd:
        filters['created_at__gte'] = sd
    if ed:
        filters['created_at__lte'] = ed

    qs = (
        Incident.objects
        .filter(**filters)
        .annotate(date=TruncDate('created_at'))
        .values('date')
        .annotate(count=Count('id'))
        .order_by('date')
    )
    return [
        {'date': str(i['date']), 'count': i['count']}
        for i in qs
    ]


def get_filtered_incidents(filters):
    qs = Incident.objects.select_related(
        'reported_by', 'assigned_to', 'area', 'machine', 'type'
    )
    area = filters.get('area')
    type_ = filters.get('type')
    status = filters.get('status')
    priority = filters.get('priority')
    start_date = filters.get('start_date')
    end_date = filters.get('end_date')
    if area:
        qs = qs.filter(area_id=area)
    if type_:
        qs = qs.filter(type_id=type_)
    if status:
        qs = qs.filter(status=status)
    if priority:
        qs = qs.filter(priority=priority)
    sd = _parse_date(start_date)
    ed = _parse_date(end_date)
    if sd:
        qs = qs.filter(created_at__gte=sd)
    if ed:
        qs = qs.filter(created_at__lte=ed)
    return qs


def build_report_data(source, filters, start_date, end_date):
    data = {}
    if source == 'metrics':
        data['avg_response_time'] = get_avg_response_time(start_date, end_date)
        data['avg_resolution_time'] = get_avg_resolution_time(start_date, end_date)
        data['resolution_rate'] = get_resolution_rate(start_date, end_date)
        data['incidents_by_area'] = get_incidents_by_area(start_date, end_date)
        data['root_cause_frequency'] = get_root_cause_frequency(start_date, end_date)
        data['critical_incidents_over_time'] = get_critical_incidents_over_time(start_date, end_date)
    elif source in ('history', 'filtered'):
        qs = get_filtered_incidents(filters) if source == 'filtered' else get_filtered_incidents({})
        data['incidents'] = list(qs.values(
            'id', 'title', 'status', 'priority', 'area__name', 'type__name',
            'reported_by__first_name', 'reported_by__last_name',
            'created_at', 'resolved_at',
        ))
    return data


def generate_report_file(report_data, file_format):
    if file_format == 'excel':
        return _generate_excel(report_data)
    return _generate_pdf(report_data)


def _generate_excel(report_data):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte'

    header_font = Font(bold=True)

    if 'incidents' in report_data:
        headers = ['ID', 'Titulo', 'Estado', 'Prioridad', 'Area', 'Tipo', 'Reportado por', 'Creado', 'Resuelto']
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
        for inc in report_data['incidents']:
            ws.append([
                inc['id'],
                inc['title'],
                inc['status'],
                inc['priority'],
                inc['area__name'],
                inc['type__name'],
                f"{inc['reported_by__first_name']} {inc['reported_by__last_name']}",
                str(inc['created_at']),
                str(inc['resolved_at'] or ''),
            ])
    elif 'avg_response_time' in report_data:
        row = 1
        ws.cell(row=row, column=1, value='Metricas').font = header_font
        row += 2

        ws.cell(row=row, column=1, value='Tiempo promedio de respuesta').font = header_font
        row += 1
        ws.cell(row=row, column=1, value='Area').font = header_font
        ws.cell(row=row, column=2, value='Minutos promedio').font = header_font
        row += 1
        for item in report_data['avg_response_time']:
            ws.cell(row=row, column=1, value=item['area_name'])
            ws.cell(row=row, column=2, value=item['avg_minutes'])
            row += 1

        row += 1
        ws.cell(row=row, column=1, value='Tiempo promedio de resolucion').font = header_font
        row += 1
        ws.cell(row=row, column=1, value='Area').font = header_font
        ws.cell(row=row, column=2, value='Minutos promedio').font = header_font
        row += 1
        for item in report_data['avg_resolution_time']:
            ws.cell(row=row, column=1, value=item['area_name'])
            ws.cell(row=row, column=2, value=item['avg_minutes'])
            row += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _generate_pdf(report_data):
    lines = ['=== Reporte de Incidencias OpsCore ===', '']
    if 'incidents' in report_data:
        lines.append(f'Total incidentes: {len(report_data["incidents"])}')
        lines.append('')
        for inc in report_data['incidents']:
            lines.append(f"#{inc['id']} - {inc['title']} [{inc['status']}]")
            lines.append(f"  Area: {inc['area__name']} | Tipo: {inc['type__name']}")
            lines.append(f"  Prioridad: {inc['priority']}")
            lines.append('')
    elif 'avg_response_time' in report_data:
        lines.append('=== Tiempo promedio de respuesta ===')
        for item in report_data['avg_response_time']:
            lines.append(f"  {item['area_name']}: {item['avg_minutes']} min")

        lines.append('')
        lines.append('=== Tiempo promedio de resolucion ===')
        for item in report_data['avg_resolution_time']:
            lines.append(f"  {item['area_name']}: {item['avg_minutes']} min")

        lines.append('')
        rr = report_data['resolution_rate']
        lines.append(f"=== Tasa de resolucion: {rr['rate']}% ===")
        lines.append(f"  Total: {rr['total']} | Cerrados: {rr['closed']}")

        lines.append('')
        lines.append('=== Incidentes por area ===')
        for item in report_data['incidents_by_area']:
            lines.append(f"  {item['area_name']}: {item['count']}")

        lines.append('')
        lines.append('=== Causa raiz frecuentes ===')
        for item in report_data['root_cause_frequency']:
            lines.append(f"  {item['type_name']} - {item['root_cause']}: {item['count']}")

        lines.append('')
        lines.append('=== Criticos en el tiempo ===')
        for item in report_data['critical_incidents_over_time']:
            lines.append(f"  {item['date']}: {item['count']}")

    content = '\n'.join(lines)
    return io.BytesIO(content.encode('utf-8'))
